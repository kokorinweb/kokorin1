"""Сквозной тест конвейера без сети: источник-заглушка + отключённые сетевые проверки.

Проверяет реальную цепочку: сбор -> нормализация -> отбраковка -> дедупликация ->
проверка -> confidence -> lead score -> SQLite -> XLSX/CSV.
"""
import asyncio

import pytest
from openpyxl import load_workbook

from config import Config
from database.models import NO_WEBSITE_HIGH_CONFIDENCE, HAS_WEBSITE
from database.repository import Repository
from exporters.csv_exporter import export_csv
from exporters.excel_exporter import export_excel
from services import pipeline as pipeline_module
from services.pipeline import Pipeline, raw_to_place
from services.social_checker import SocialOutcome, SocialProfile
from services.website_checker import DomainCheck
from sources import registry
from sources.base import BaseSource, RawPlace
from utils.cities import find_city


class StubSource(BaseSource):
    """Отдаёт фиксированный набор записей — сети не касается."""

    name = "stub"
    priority = 99
    PLACES = [
        RawPlace(source="stub", source_id="1", name="Ресторан «Белый Кролик»",
                 raw_category="restaurant", city="Казань", region="Республика Татарстан",
                 address="ул. Баумана, 15", latitude=55.7900, longitude=49.1220,
                 phones=["+7 (843) 555-11-22"], rating=4.6, reviews_count=320,
                 working_hours="пн-вс 10:00-23:00",
                 social_urls=["https://vk.com/belyykrolik"], source_url="https://example.org/1"),
        # дубль первой записи из другого «источника»
        RawPlace(source="stub", source_id="2", name="Белый Кролик",
                 raw_category="ресторан", city="Казань", region="Республика Татарстан",
                 address="г. Казань, улица Баумана, д. 15", latitude=55.7901, longitude=49.1221,
                 phones=["88435551122"], rating=4.5, reviews_count=210),
        # у этого сайт указан в карточке
        RawPlace(source="stub", source_id="3", name="Кафе Мимими",
                 raw_category="cafe", city="Казань", region="Республика Татарстан",
                 address="ул. Кремлёвская, 8", latitude=55.7960, longitude=49.1080,
                 phones=["+79170001122"], website="https://mimimi-kazan.ru",
                 rating=4.2, reviews_count=95),
        # мусор — должен отсеяться
        RawPlace(source="stub", source_id="4", name="Столовая №12",
                 raw_category="столовая", city="Казань", region="Республика Татарстан",
                 address="ул. Заводская, 1"),
        RawPlace(source="stub", source_id="5", name="Кафе",
                 raw_category="cafe", city="Казань", region="Республика Татарстан"),
        # другой филиал — сливать нельзя
        RawPlace(source="stub", source_id="6", name="Ресторан «Белый Кролик»",
                 raw_category="restaurant", city="Казань", region="Республика Татарстан",
                 address="пр. Победы, 141", latitude=55.7500, longitude=49.2000,
                 phones=["+7 843 555-99-88"], rating=4.4, reviews_count=60),
    ]

    async def fetch_city(self, city, categories, limit):
        return list(self.PLACES)


@pytest.fixture
def config(tmp_path):
    cfg = Config()
    cfg.db_path = str(tmp_path / "test.sqlite3")
    cfg.sources = ["stub"]
    cfg.use_search = False          # без поискового API
    cfg.use_socials = True          # соцсети проверяем (сетевой вызов замокан)
    cfg.use_domain_probe = True     # перебор доменов оставляем, но HTTP замокан
    cfg.progress_bar = False
    cfg.limit = 100
    cfg.min_confidence = 85
    return cfg


@pytest.fixture
def patched(monkeypatch):
    monkeypatch.setitem(registry.SOURCE_CLASSES, "stub", StubSource)

    async def fake_check_domain(self, domain, *, name, phones, city, address):
        """mimimi-kazan.ru — живой сайт заведения, всё остальное не существует."""
        if domain == "mimimi-kazan.ru":
            return DomainCheck(
                domain=domain, url=f"https://{domain}/", reachable=True, http_status=200,
                final_url=f"https://{domain}/", final_domain=domain,
                url_type="OFFICIAL_WEBSITE", title="Кафе Мимими — Казань",
                name_match=100.0, phone_match=True, city_match=True, score=95,
                verified=True, reason="название 100%, телефон совпал",
            )
        return DomainCheck(domain=domain, reachable=False, error="домен не существует",
                           reason="домен не существует")

    monkeypatch.setattr(
        "services.website_checker.WebsiteChecker.check_domain", fake_check_domain
    )

    async def fake_social_check(self, *, name, city, social_urls, search_socials=None):
        """Как при настроенном VK-токене: профиль прочитан, ссылки на сайт в нём нет."""
        outcome = SocialOutcome(performed=True)
        for url in social_urls or []:
            outcome.profiles.append(
                SocialProfile(network="vk", url=url, handle="x", found=True,
                              inspected=True, followers=4200)
            )
        outcome.has_any_profile = bool(outcome.profiles)
        outcome.total_followers = sum(p.followers or 0 for p in outcome.profiles)
        return outcome

    monkeypatch.setattr("services.social_checker.SocialChecker.check", fake_social_check)
    return True


def test_raw_to_place_filters_and_normalizes():
    place = raw_to_place(StubSource.PLACES[0])
    assert place is not None
    assert place.phone == "+78435551122"
    assert place.category == "restaurant"
    assert place.vk_url == "https://vk.com/belyykrolik"
    assert raw_to_place(StubSource.PLACES[3]) is None      # столовая
    assert raw_to_place(StubSource.PLACES[4]) is None      # безымянное «Кафе»


def test_social_link_in_website_field_is_not_a_website():
    raw = RawPlace(source="stub", source_id="x", name="Бар Лаунж", raw_category="bar",
                   city="Казань", website="https://vk.com/barlounge")
    place = raw_to_place(raw)
    assert place.website_from_source == ""
    assert place.vk_url == "https://vk.com/barlounge"


def test_full_pipeline(config, patched, tmp_path):
    repo = Repository(config.db_path)
    pipeline = Pipeline(config, repo)
    city = find_city("Казань")

    stats = asyncio.run(pipeline.run_city(city, 100))
    asyncio.run(pipeline.aclose())

    # 6 сырых -> 2 отброшено -> 4 записи -> 1 дубль слит -> 3 уникальных
    assert stats.found == 4
    assert stats.after_dedup == 3
    assert stats.checked == 3

    places = repo.list_places()
    assert len(places) == 3
    by_name = {p.name: p for p in places}

    # филиал на Баумана — с профилем VK, проверенным насквозь
    krolik = next(p for p in places if "Баумана" in p.full_address)
    assert krolik.website_status == NO_WEBSITE_HIGH_CONFIDENCE
    assert krolik.website_confidence >= 85
    assert krolik.lead_score > 0
    assert "+78435551122" in krolik.phones

    mimimi = by_name["Кафе Мимими"]
    assert mimimi.website_status == HAS_WEBSITE
    assert mimimi.website_confidence == 0
    assert mimimi.official_website.startswith("https://mimimi-kazan.ru")

    # два филиала «Белого Кролика» остались раздельными записями
    assert sum(1 for p in places if "Кролик" in p.name) == 2

    # --- экспорт ---
    xlsx = tmp_path / "leads.xlsx"
    csv = tmp_path / "leads.csv"
    export_excel(places, xlsx)
    export_csv([p for p in places if p.website_confidence >= 85], csv)

    assert xlsx.exists() and csv.exists()
    book = load_workbook(xlsx)
    assert book.sheetnames == ["All leads", "High confidence", "Needs review",
                               "Has website", "Errors"]
    sheet = book["All leads"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    assert sheet.max_row == 4              # шапка + 3 записи

    high = book["High confidence"]
    assert high.max_row >= 2

    # сортировка по lead_score DESC
    header = [c.value for c in sheet[1]]
    score_col = header.index("Lead score") + 1
    scores = [sheet.cell(row=r, column=score_col).value for r in range(2, sheet.max_row + 1)]
    assert scores == sorted(scores, reverse=True)

    repo.close()


def test_resume_skips_already_checked(config, patched):
    repo = Repository(config.db_path)
    city = find_city("Казань")

    first = Pipeline(config, repo)
    asyncio.run(first.run_city(city, 100))
    asyncio.run(first.aclose())

    second = Pipeline(config, repo)
    stats = asyncio.run(second.run_city(city, 100))
    asyncio.run(second.aclose())

    # повторный прогон не перепроверяет то, что проверено недавно
    assert stats.checked == 0
    assert stats.from_cache == 3
    repo.close()


def test_cache_and_progress_survive_restart(config, patched):
    repo = Repository(config.db_path)
    repo.cache_set("search", "k", {"hits": []}, 3600)
    repo.mark_progress("stub", "Казань", "done", found=4)
    repo.close()

    repo2 = Repository(config.db_path)
    assert repo2.cache_get("search", "k") == {"hits": []}
    assert repo2.progress_status("stub", "Казань") == "done"
    repo2.close()
