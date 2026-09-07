"""Экспорт в XLSX: несколько листов, фильтры, закреплённая шапка, ссылки."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from database.models import Place
from exporters.common import COLUMNS, LINK_COLUMNS, SHEETS, filter_places, places_to_frame
from utils.logger import get_logger

log = get_logger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")
MAX_LINK_ROWS = 20_000        # гиперссылки дорогие, на огромных выгрузках не ставим


def _style_sheet(sheet: Worksheet, frame: pd.DataFrame) -> None:
    if frame.empty:
        sheet.cell(row=1, column=1, value="нет записей")
        return

    widths = {title: width for _, title, width, _ in COLUMNS}
    headers = list(frame.columns)

    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(title, 18)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(frame) + 1}"
    )
    sheet.row_dimensions[1].height = 30

    if len(frame) <= MAX_LINK_ROWS:
        link_indexes = [i for i, title in enumerate(headers, start=1) if title in LINK_COLUMNS]
        for row_index in range(2, len(frame) + 2):
            for col_index in link_indexes:
                cell = sheet.cell(row=row_index, column=col_index)
                value = cell.value
                if isinstance(value, str) and value.startswith(("http://", "https://")):
                    cell.hyperlink = value
                    cell.font = LINK_FONT

    # подсветка confidence и lead score
    for title, rule in (
        ("Confidence «нет сайта»", CellIsRule(operator="greaterThanOrEqual", formula=["85"],
                                              fill=PatternFill("solid", fgColor="C6EFCE"))),
        ("Lead score", CellIsRule(operator="greaterThanOrEqual", formula=["70"],
                                  fill=PatternFill("solid", fgColor="FFF2CC"))),
    ):
        if title in headers:
            letter = get_column_letter(headers.index(title) + 1)
            sheet.conditional_formatting.add(f"{letter}2:{letter}{len(frame) + 1}", rule)


def export_excel(places: list[Place], path: str | Path, *, top: int | None = None) -> Path:
    """Пишет XLSX со всеми листами. ``top`` ограничивает лист All leads."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        for spec in SHEETS:
            subset = filter_places(places, spec)
            frame = places_to_frame(subset)
            if spec.name == "All leads" and top:
                frame = frame.head(top)
            frame.to_excel(writer, sheet_name=spec.name[:31], index=False)
            _style_sheet(writer.sheets[spec.name[:31]], frame)

    log.info("XLSX сохранён: %s", target)
    return target
